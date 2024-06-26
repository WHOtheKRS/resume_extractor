import re
import PyPDF2
from openpyxl import Workbook
import os

def extract_text_from_pdf(pdf_file):
    
    with open(pdf_file, 'rb') as f:
        pdf_reader = PyPDF2.PdfReader(f)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
    return text.strip()

def extract_text_from_docx(docx_file):
    
    document = Document(docx_file)
    text = ""
    for paragraph in document.paragraphs:
        text += paragraph.text
    return text.strip()

def extract_info(filename):
    """
    Extracts email, phone number, and overall text from a PDF resume.
    """
    with open(filename, 'rb') as f:
        pdf_reader = PyPDF2.PdfReader(f)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()

    email = re.findall(r"[\w\.-]+@[\w\.-]+\.[\w]{2,}", text)
    phone_number = re.findall(r"\d{3}-\d{3}-\d{4}|\(\d{3}\) \d{3}-\d{4}", text)
    return {
        "filename": filename,
        "email": email[0] if email else "",  # Assuming only one email exists
        "phone_number": phone_number[0] if phone_number else "",  # Assuming one number
        "text": text
    }

def create_xlsx(resume_data, output_filename):
    
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Filename", "Email ID", "Contact Number", "Text"])

    for resume in resume_data:
        worksheet.append([resume["filename"], resume["email"], resume["phone_number"], resume["text"]])

    workbook.save(output_filename)

def process_resumes(resume_dir):
   
    resume_data = []
    for filename in os.listdir(resume_dir):
        if filename.endswith(".pdf"):
            resume_data.append(extract_info(os.path.join(resume_dir, filename)))
        else:
            print(f"Skipping unsupported file: {filename}")

    output_filename = os.path.join(resume_dir, "extracted_data.xlsx")
    create_xlsx(resume_data, output_filename)
